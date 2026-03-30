"""FAQ loader — reads FAQ.xlsx and transforms rows into ChromaDB documents."""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Any, Dict, List

import openpyxl

logger = logging.getLogger(__name__)

# Normalize regime names → internal keys
REGIME_NORMALIZE: Dict[str, str] = {
    "бюджетное планирование": "budget_planning",
    "исполнение бюджета": "budget_execution",
    "медиамониторинг": "media_monitoring",
    "личный кабинет": "personal_account",
}

# Normalize module names → internal keys
MODULE_NORMALIZE: Dict[str, str] = {
    "бюджетные запросы": "budget_requests",
    "формирование": "formation",
    "заявки на внесение изменений": "change_requests",
    "мониторинг исполнения": "execution_monitoring",
    "аналитика": "analytics",
    "подписанты": "signatories",
    "аналитика исполнения": "execution_analytics",
    "лимиты": "limits",
}

# Normalize topic names → internal keys
TOPIC_NORMALIZE: Dict[str, str] = {
    "свод по абп/гу/гккп": "svod_abp_gu_gkkp",
    "форма расчетов": "calculation_forms",
    "штатное расписание": "staffing",
    "форма гу": "forma_gu",
    "индивидуальный план финансирования": "ipf",
    "заявки на внесение изменений": "change_requests",
    "заявки гу": "gu_requests",
    "план финансирования": "finance_plan",
    "текущие расходы": "current_expenses",
    "лимиты гу": "limits_gu",
    "настройки гккп": "gkkp_settings",
    "аналитика планирования": "planning_analytics",
    "пакетная выгрузка": "batch_export",
    "отчеты": "reports",
}


def _normalize(value: str, mapping: Dict[str, str]) -> str:
    """Normalize a Russian string to an internal key."""
    if not value:
        return ""
    lower = value.strip().lower()
    return mapping.get(lower, re.sub(r"\s+", "_", lower))


def _derive_source_page(regime: str, module: str, topic: str) -> str:
    """Derive a source_page identifier for Module 3 context filtering."""
    if topic:
        return f"page_{regime.split('_')[0]}_{topic}"
    if module:
        return f"page_{regime.split('_')[0]}_{module}"
    return f"page_{regime}"


def load_faq(faq_path: str) -> List[Dict[str, Any]]:
    """Load FAQ.xlsx and return ChromaDB-ready documents.

    Expected columns: Title, [Answer], Topic, Module, Regime, URL
    The Answer column is optional — if missing, content = question only.
    """
    path = Path(faq_path)
    if not path.exists():
        logger.error("FAQ file not found: %s", faq_path)
        return []

    wb = openpyxl.load_workbook(str(path), read_only=True)
    ws = wb.active
    if ws is None:
        logger.error("No active sheet in %s", faq_path)
        return []

    # Detect columns
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    headers_lower = [h.lower().strip() if h else "" for h in headers]

    col_map = {}
    for i, h in enumerate(headers_lower):
        if "title" in h or "вопрос" in h:
            col_map["title"] = i
        elif "answer" in h or "ответ" in h:
            col_map["answer"] = i
        elif "topic" in h or "тема" in h or "раздел" in h:
            col_map["topic"] = i
        elif "module" in h or "модуль" in h:
            col_map["module"] = i
        elif "regime" in h or "режим" in h:
            col_map["regime"] = i
        elif "url" in h or "ссылка" in h:
            col_map["url"] = i

    if "title" not in col_map:
        logger.error("FAQ file missing 'Title' column. Headers: %s", headers)
        return []

    has_answer = "answer" in col_map
    if not has_answer:
        logger.warning("FAQ file has no 'Answer' column — documents will contain questions only")

    documents: List[Dict[str, Any]] = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        row_list = list(row)

        title = row_list[col_map["title"]] if col_map.get("title") is not None else None
        if not title:
            continue

        # Clean numbered prefix: "1. Как..." → "Как..."
        question = re.sub(r"^\d+\.\s*", "", str(title).strip())

        answer = ""
        if has_answer and col_map.get("answer") is not None:
            answer = str(row_list[col_map["answer"]] or "").strip()

        topic_raw = str(row_list[col_map.get("topic", 0)] or "").strip() if "topic" in col_map else ""
        module_raw = str(row_list[col_map.get("module", 0)] or "").strip() if "module" in col_map else ""
        regime_raw = str(row_list[col_map.get("regime", 0)] or "").strip() if "regime" in col_map else ""
        url = str(row_list[col_map.get("url", 0)] or "").strip() if "url" in col_map else ""

        # Normalize
        topic = _normalize(topic_raw, TOPIC_NORMALIZE)
        module = _normalize(module_raw, MODULE_NORMALIZE)
        regime = _normalize(regime_raw, REGIME_NORMALIZE)
        source_page = _derive_source_page(regime, module, topic)

        # Build content
        if answer:
            content = f"Вопрос: {question}\n\nОтвет: {answer}"
        else:
            content = f"Вопрос: {question}"

        doc_id = f"faq_{row_idx:04d}"
        documents.append(
            {
                "id": doc_id,
                "content": content,
                "metadata": {
                    "type": "faq",
                    "topic": topic,
                    "module": module,
                    "regime": regime,
                    "source_page": source_page,
                    "language": "ru",
                    "video_url": url,
                    "image_url": "",
                    "original_question": question,
                },
            }
        )

    wb.close()
    logger.info("Loaded %d FAQ documents from %s", len(documents), faq_path)
    return documents
